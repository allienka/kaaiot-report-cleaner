from pathlib import Path
from shutil import move
from openpyxl import load_workbook
from datetime import datetime

# Define folders
incoming_dir = Path("incoming")
output_dir = Path("output")
processed_dir = Path("processed")

# Create folders if missing
output_dir.mkdir(exist_ok=True)
processed_dir.mkdir(exist_ok=True)

# Find all Excel files
excel_files = incoming_dir.glob("*.xlsx")

for input_file in excel_files:

    print(f"Processing: {input_file.name}")

    # Create output filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    output_file = (
    output_dir /
    f"{input_file.stem}_fixed_{timestamp}.xlsx"
    )

    try:
        # Load workbook
        wb = load_workbook(input_file)

        # Process all sheets
        for ws in wb.worksheets:

            for row in ws.iter_rows():
                for cell in row:

                    # Only process text cells
                    if isinstance(cell.value, str):

                        value = cell.value.strip()

                        try:
                            # Decimal numbers
                            if "." in value:
                                number = float(value)

                            # Integers
                            else:
                                number = int(value)

                            # Replace text with number
                            cell.value = number

                            print(
                                f"Converted {ws.title} "
                                f"{cell.coordinate}: "
                                f"{value} -> {number}"
                            )

                        except ValueError:
                            pass

        # Save cleaned workbook
        wb.save(output_file)

        print(f"Saved cleaned file: {output_file}")

        # Move original file to processed folder
        move(str(input_file), processed_dir / input_file.name)

        print(f"Moved original to processed/")

    except Exception as e:
        print(f"Error processing {input_file.name}: {e}")