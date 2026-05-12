import sys
from pathlib import Path
from openpyxl import load_workbook


# Check command line arguments
if len(sys.argv) != 3:
    print("Usage: python fix_report.py <input_file> <output_file>")
    sys.exit(1)

# Get file paths from arguments
input_file = Path(sys.argv[1])
output_file = Path(sys.argv[2])

# Load workbook
wb = load_workbook(input_file)

# Loop through all sheets
for ws in wb.worksheets:

    # Loop through all rows and cells
    for row in ws.iter_rows():
        for cell in row:

            # Process only text cells
            if isinstance(cell.value, str):

                value = cell.value.strip()

                try:
                    # Convert decimal numbers
                    if "." in value:
                        number = float(value)

                    # Convert integers
                    else:
                        number = int(value)

                    # Replace text with real number
                    cell.value = number
                    
                    print(f"Converted {cell.coordinate}: {value} -> {number}")

                except ValueError:
                    # Ignore non-numeric text
                    pass

# Save corrected workbook
wb.save(output_file)

print(f"Saved corrected file: {output_file}")